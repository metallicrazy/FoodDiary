#!/usr/bin/env python3
"""
Convert the GOV.UK CoFID spreadsheet (McCance & Widdowson's Composition of Foods
Integrated Dataset) into data/cofid.json so FoodDiary ships with offline UK foods.

Optional — the app can also import the spreadsheet directly from Me ▸ UK food database.
Use this script if you'd rather bundle the data in the repo so every device gets it.

Usage:
    pip install openpyxl
    python tools/build_cofid.py "McCance_Widdowsons_Composition_of_Foods_Integrated_Dataset_2021.xlsx"

Download the spreadsheet from:
    https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid
Data © Crown copyright, Open Government Licence v3.
"""
import json, re, sys, os

try:
    import openpyxl
except ImportError:
    sys.exit("Please install openpyxl first:  pip install openpyxl")

if len(sys.argv) < 2:
    sys.exit(__doc__)

path = sys.argv[1]
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
sheet_name = next((n for n in wb.sheetnames if re.search('proximates', n, re.I)), wb.sheetnames[0])
ws = wb[sheet_name]
rows = [list(r) for r in ws.iter_rows(values_only=True)]

low = lambda v: str(v or '').strip().lower()
header_idx = next((i for i, r in enumerate(rows[:6]) if any(low(c) == 'food code' for c in r)), -1)
abbr_idx = next((i for i, r in enumerate(rows[:6]) if any(low(c) in ('kcals', 'prot') for c in r)), -1)
header = [low(c) for c in rows[header_idx]] if header_idx >= 0 else []
abbr = [low(c) for c in rows[abbr_idx]] if abbr_idx >= 0 else []

def find(abbrevs, name_re):
    for a in abbrevs:
        if a in abbr:
            return abbr.index(a)
    for i, h in enumerate(header):
        if re.search(name_re, h):
            return i
    return -1

c_code = find(['food code'], r'^food code')
c_name = find(['name', 'food name'], r'^food name')
c_group = find(['group'], r'^group$')
c_prot = find(['prot'], r'^protein')
c_fat = find(['fat'], r'^fat\b')
c_cho = find(['cho'], r'^carbohydrate')
c_kcal = find(['kcals'], r'energy.*kcal')
c_kj = find(['kj'], r'energy.*kj')
c_sug = find(['totsug'], r'^total sugars')
c_fib = find(['aoacfib'], r'aoac')
c_sat = find(['satfod'], r'^satd fa.*food|^saturated.*food')

def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ('', 'N'):
        return None
    if s == 'Tr':
        return 0.0
    s = re.sub(r'[^0-9.\-]', '', s)
    try:
        return float(s)
    except ValueError:
        return None

def get(row, i):
    return row[i] if 0 <= i < len(row) else None

start = max(header_idx, abbr_idx) + 1
out = []
for r in rows[start:]:
    name = str(get(r, c_name) or '').strip()
    if not name:
        continue
    kcal = num(get(r, c_kcal))
    if kcal is None and c_kj >= 0:
        kj = num(get(r, c_kj))
        kcal = None if kj is None else kj / 4.184
    if kcal is None:
        continue
    food = {
        'id': 'cofid:' + str(get(r, c_code) or len(out)),
        'source': 'cofid',
        'code': str(get(r, c_code) or ''),
        'name': name,
        'group': str(get(r, c_group) or ''),
        'kcal': round(kcal),
        'protein': num(get(r, c_prot)) or 0,
        'carbs': num(get(r, c_cho)) or 0,
        'fat': num(get(r, c_fat)) or 0,
    }
    for key, col in (('sugars', c_sug), ('fibre', c_fib), ('satFat', c_sat)):
        if col >= 0:
            v = num(get(r, col))
            if v is not None:
                food[key] = v
    out.append(food)

os.makedirs('data', exist_ok=True)
with open('data/cofid.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, separators=(',', ':'))
with open('data/cofid.js', 'w', encoding='utf-8') as f:  # same data as a script, for local file:// preview
    f.write('window.__COFID=' + json.dumps(out, ensure_ascii=False, separators=(',', ':')) + ';')
print(f'Wrote {len(out)} foods from sheet "{sheet_name}" to data/cofid.json and data/cofid.js')
